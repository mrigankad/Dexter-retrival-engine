"""
Word Builder — Professional Word Document Generation.

Uses python-docx to create styled Word documents with
headers, tables, charts, and source citations.
"""

import io
import re
from typing import List, Dict, Any, Optional
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE

import pandas as pd


# ── Brand Colors ─────────────────────────────────────────────
BRAND_PRIMARY = RGBColor(0x3C, 0x8C, 0x81)
BRAND_DARK = RGBColor(0x2D, 0x6A, 0x5F)
BRAND_TEXT = RGBColor(0x2C, 0x3E, 0x50)
BRAND_MUTED = RGBColor(0x7F, 0x8C, 0x8D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


class WordBuilder:
    """Builds professional Word documents."""

    def __init__(self, title: str = "RAG Analysis Report"):
        self.doc = Document()
        self.title = title
        self._setup_styles()

    def _setup_styles(self):
        """Configure document-wide styles."""
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10)
        style.font.color.rgb = BRAND_TEXT

        # Update heading styles
        for level in range(1, 4):
            heading_style = self.doc.styles[f"Heading {level}"]
            heading_style.font.color.rgb = BRAND_PRIMARY
            heading_style.font.name = "Calibri"

    def add_header(self, query: str = ""):
        """Add report header."""
        # Title
        title_para = self.doc.add_heading(self.title, level=0)
        for run in title_para.runs:
            run.font.color.rgb = BRAND_TEXT
            run.font.size = Pt(22)

        # Subtitle
        subtitle = self.doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = subtitle.add_run(
            f"RAG-Powered Analysis Report  |  {datetime.now().strftime('%B %d, %Y')}"
        )
        run.font.size = Pt(11)
        run.font.color.rgb = BRAND_MUTED

        # Horizontal rule
        self.doc.add_paragraph("_" * 80)

        # Query box
        if query:
            query_para = self.doc.add_paragraph()
            run_label = query_para.add_run("Query: ")
            run_label.bold = True
            run_label.font.color.rgb = BRAND_PRIMARY
            run_text = query_para.add_run(query)
            run_text.font.color.rgb = BRAND_TEXT

        self.doc.add_paragraph()

    def add_section(self, title: str, content: str, level: int = 1):
        """Add a titled section with content."""
        if title:
            self.doc.add_heading(title, level=level)

        for para_text in content.split("\n"):
            para_text = para_text.strip()
            if not para_text:
                continue

            # Bullet points
            if para_text.startswith(("- ", "• ", "* ")):
                text = para_text[2:].strip()
                p = self.doc.add_paragraph(style="List Bullet")
                self._add_formatted_run(p, text)
            # Numbered lists
            elif re.match(r"^\d+\.", para_text):
                p = self.doc.add_paragraph(style="List Number")
                self._add_formatted_run(p, para_text)
            else:
                p = self.doc.add_paragraph()
                self._add_formatted_run(p, para_text)

    def add_answer(self, answer: str, section_title: str = "Detailed Analysis"):
        """Add the RAG answer with intelligent formatting."""
        self.doc.add_heading(section_title, level=1)

        lines = answer.split("\n")
        for line in lines:
            line = line.strip()
            if not line:
                continue

            # Headers
            if line.startswith("### "):
                self.doc.add_heading(line[4:], level=3)
            elif line.startswith("## "):
                self.doc.add_heading(line[3:], level=2)
            elif line.startswith("**") and line.endswith("**"):
                p = self.doc.add_paragraph()
                run = p.add_run(line.strip("*"))
                run.bold = True
                run.font.color.rgb = BRAND_DARK
            # Bullet points
            elif line.startswith(("- ", "• ", "* ")):
                text = line[2:].strip()
                p = self.doc.add_paragraph(style="List Bullet")
                self._add_formatted_run(p, text)
            # Numbered points
            elif re.match(r"^\d+\.", line):
                p = self.doc.add_paragraph(style="List Number")
                self._add_formatted_run(p, line)
            else:
                p = self.doc.add_paragraph()
                self._add_formatted_run(p, line)

    def add_kpi_table(self, kpis: Dict[str, Any], title: str = "Key Metrics"):
        """Add KPIs as a styled table."""
        if not kpis:
            return

        self.doc.add_heading(title, level=1)

        items = list(kpis.items())
        cols = min(len(items), 3)
        rows_needed = (len(items) + cols - 1) // cols

        table = self.doc.add_table(rows=rows_needed * 2, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, (key, val) in enumerate(items):
            row = (i // cols) * 2
            col = i % cols

            # Value cell
            val_str = f"{val:,.1f}" if isinstance(val, float) and val != int(val) else str(val)
            val_cell = table.cell(row, col)
            val_para = val_cell.paragraphs[0]
            val_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = val_para.add_run(val_str)
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = BRAND_PRIMARY

            # Label cell
            lbl_cell = table.cell(row + 1, col)
            lbl_para = lbl_cell.paragraphs[0]
            lbl_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = lbl_para.add_run(key)
            run.font.size = Pt(9)
            run.font.color.rgb = BRAND_MUTED

        self.doc.add_paragraph()

    def add_chart_image(self, chart_bytes: bytes, caption: str = "", width: float = 5.5):
        """Add a chart image to the document."""
        img_stream = io.BytesIO(chart_bytes)
        self.doc.add_picture(img_stream, width=Inches(width))

        last_para = self.doc.paragraphs[-1]
        last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if caption:
            cap_para = self.doc.add_paragraph()
            cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap_para.add_run(caption)
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = BRAND_MUTED

    def add_table(self, df: pd.DataFrame, title: str = ""):
        """Add a pandas DataFrame as a styled table."""
        if title:
            self.doc.add_heading(title, level=2)

        rows = len(df) + 1  # +1 for header
        cols = len(df.columns)
        table = self.doc.add_table(rows=rows, cols=cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Headers
        for j, col_name in enumerate(df.columns):
            cell = table.cell(0, j)
            cell.text = str(col_name)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.color.rgb = WHITE
                    run.font.size = Pt(10)
            # Header background color
            shading = cell._element
            from docx.oxml.ns import qn
            tcPr = shading.get_or_add_tcPr()
            shd = tcPr.makeelement(qn("w:shd"), {
                qn("w:fill"): "3C8C81",
                qn("w:val"): "clear",
            })
            tcPr.append(shd)

        # Data rows
        for i, (_, row) in enumerate(df.iterrows()):
            for j, val in enumerate(row):
                cell = table.cell(i + 1, j)
                cell.text = str(val) if val is not None else ""
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.font.color.rgb = BRAND_TEXT

        self.doc.add_paragraph()

    def add_sources(self, sources: List[dict], title: str = "Source References"):
        """Add source references."""
        self.doc.add_heading(title, level=1)

        for i, src in enumerate(sources, 1):
            p = self.doc.add_paragraph()
            run = p.add_run(f"[{i}] ")
            run.bold = True
            run.font.color.rgb = BRAND_PRIMARY

            page = src.get("page", "N/A")
            source = src.get("source", "Unknown")
            score = src.get("score", 0)
            run = p.add_run(f"Page {page} | {source} | Score: {score:.3f}")
            run.font.size = Pt(9)
            run.font.color.rgb = BRAND_MUTED

            content = src.get("content", "")
            if content:
                preview = content[:200] + "..." if len(content) > 200 else content
                p2 = self.doc.add_paragraph()
                run = p2.add_run(preview)
                run.font.size = Pt(9)
                run.font.color.rgb = BRAND_TEXT
                run.italic = True

    def add_page_break(self):
        self.doc.add_page_break()

    def build(self, output_path: str = None) -> bytes:
        """Build the Word document and return as bytes."""
        buf = io.BytesIO()
        self.doc.save(buf)
        buf.seek(0)
        data = buf.getvalue()

        if output_path:
            with open(output_path, "wb") as f:
                f.write(data)

        return data

    # ── Helpers ───────────────────────────────────────────────

    def _add_formatted_run(self, paragraph, text: str):
        """Add text with basic markdown bold formatting."""
        text = text.replace("**", "%%BOLD%%")
        parts = text.split("%%BOLD%%")
        for i, part in enumerate(parts):
            if not part:
                continue
            run = paragraph.add_run(part)
            run.font.size = Pt(10)
            run.font.color.rgb = BRAND_TEXT
            if i % 2 == 1:  # Odd sections are bold
                run.bold = True


# ── Convenience Function ─────────────────────────────────────

def generate_word_report(
    query: str,
    answer: str,
    sources: List[dict] = None,
    charts: List[tuple] = None,
    kpis: Dict[str, Any] = None,
    title: str = "RAG Analysis Report",
    output_path: str = None,
) -> bytes:
    """One-call function to generate a complete Word report."""
    doc = WordBuilder(title=title)

    doc.add_header(query=query)

    if kpis:
        doc.add_kpi_table(kpis)

    if charts:
        doc.add_section("Data Visualization", "")
        for chart_bytes, caption in charts:
            doc.add_chart_image(chart_bytes, caption=caption)

    doc.add_answer(answer)

    if sources:
        doc.add_page_break()
        doc.add_sources(sources)

    return doc.build(output_path)
