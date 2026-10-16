"""
File parsers producing ``ParsedDocument`` (page-level text, tables, headings).

PDF  : Docling (if installed and PARSER_BACKEND allows) → PyMuPDF → pypdf
DOCX : python-docx (headings become markdown headings, tables become markdown tables)
TXT/MD/HTML/CSV : stdlib
"""

from __future__ import annotations

import csv
import html
import os
import re
from typing import List, Optional

from src.contracts import ParsedDocument, ParsedPage
from src.core.config import settings
from src.core.logger import get_logger

logger = get_logger(__name__)


class ParserError(RuntimeError):
    pass


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$", re.MULTILINE)


def _extract_headings(text: str) -> List[str]:
    return [m.group(2).strip() for m in _HEADING_RE.finditer(text)]


def _extract_md_tables(text: str) -> List[str]:
    tables: List[str] = []
    block: List[str] = []
    for line in text.splitlines():
        if line.strip().startswith("|") and line.strip().endswith("|"):
            block.append(line.strip())
        else:
            if len(block) >= 2:
                tables.append("\n".join(block))
            block = []
    if len(block) >= 2:
        tables.append("\n".join(block))
    return tables


# ── PDF ─────────────────────────────────────────────────────────────────

def _parse_pdf_docling(path: str) -> Optional[ParsedDocument]:
    try:
        from docling.document_converter import DocumentConverter  # type: ignore
    except Exception:
        return None
    try:
        result = DocumentConverter().convert(path)
        doc = result.document
        pages: List[ParsedPage] = []
        # Export per page when possible; fall back to a single markdown page.
        n_pages = len(getattr(doc, "pages", {}) or {}) or 1
        if n_pages > 1 and hasattr(doc, "export_to_markdown"):
            for page_no in range(1, n_pages + 1):
                try:
                    md = doc.export_to_markdown(page_no=page_no)
                except TypeError:
                    md = ""
                if md.strip():
                    pages.append(ParsedPage(page=page_no, text=md, tables=_extract_md_tables(md), headings=_extract_headings(md)))
        if not pages:
            md = doc.export_to_markdown()
            pages = [ParsedPage(page=1, text=md, tables=_extract_md_tables(md), headings=_extract_headings(md))]
        return ParsedDocument(filename=os.path.basename(path), file_type="pdf", pages=pages, parser="docling")
    except Exception as exc:
        logger.warning("Docling failed, falling back", error=str(exc))
        return None


def _parse_pdf_pymupdf(path: str) -> Optional[ParsedDocument]:
    try:
        import fitz  # type: ignore
    except Exception:
        return None
    pages: List[ParsedPage] = []
    with fitz.open(path) as doc:
        for page in doc:
            text = page.get_text("text") or ""
            tables: List[str] = []
            try:
                for t in page.find_tables().tables:  # PyMuPDF >= 1.23
                    md = t.to_markdown()
                    if md and md.count("|") > 4:
                        tables.append(md.strip())
            except Exception:
                pass
            if text.strip() or tables:
                pages.append(ParsedPage(page=page.number + 1, text=text.strip(), tables=tables))
    return ParsedDocument(filename=os.path.basename(path), file_type="pdf", pages=pages, parser="pymupdf")


def _parse_pdf_pypdf(path: str) -> ParsedDocument:
    from pypdf import PdfReader
    reader = PdfReader(path)
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(ParsedPage(page=i, text=text))
    title = None
    try:
        title = (reader.metadata or {}).get("/Title")
    except Exception:
        pass
    return ParsedDocument(filename=os.path.basename(path), file_type="pdf", pages=pages, parser="pypdf", title=title or None)


def parse_pdf(path: str) -> ParsedDocument:
    backend = settings.parser_backend
    doc: Optional[ParsedDocument] = None
    if backend in ("auto", "docling"):
        doc = _parse_pdf_docling(path)
    if doc is None and backend in ("auto", "docling", "pymupdf"):
        doc = _parse_pdf_pymupdf(path)
    if doc is None:
        doc = _parse_pdf_pypdf(path)
    return doc


# ── DOCX ────────────────────────────────────────────────────────────────

def _docx_table_to_md(table) -> str:
    rows = []
    for r in table.rows:
        cells = [" ".join(c.text.split()) for c in r.cells]
        rows.append("| " + " | ".join(cells) + " |")
    if not rows:
        return ""
    header_sep = "| " + " | ".join("---" for _ in table.rows[0].cells) + " |"
    return "\n".join([rows[0], header_sep] + rows[1:])


def parse_docx(path: str) -> ParsedDocument:
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = DocxDocument(path)
    parts: List[str] = []
    tables: List[str] = []
    headings: List[str] = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.startswith("Heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                parts.append(f"{'#' * min(int(level), 6)} {text}")
                headings.append(text)
            elif style == "Title":
                parts.append(f"# {text}")
                headings.append(text)
            else:
                parts.append(text)
        elif tag == "tbl":
            md = _docx_table_to_md(Table(child, doc))
            if md:
                tables.append(md)
                parts.append(md)
    text = "\n\n".join(parts)
    # DOCX has no fixed pages: approximate 3000 chars per page for citations.
    pages: List[ParsedPage] = []
    page_size = 3000
    buf: List[str] = []
    size = 0
    for block in parts:
        if size + len(block) > page_size and buf:
            pages.append(ParsedPage(page=len(pages) + 1, text="\n\n".join(buf)))
            buf, size = [], 0
        buf.append(block)
        size += len(block)
    if buf:
        pages.append(ParsedPage(page=len(pages) + 1, text="\n\n".join(buf)))
    for p in pages:
        p.tables = _extract_md_tables(p.text)
        p.headings = _extract_headings(p.text)
    title = doc.core_properties.title or (headings[0] if headings else None)
    return ParsedDocument(filename=os.path.basename(path), file_type="docx", pages=pages, parser="python-docx",
                          title=title, metadata={"approx_pages": True})


# ── Plain text / markdown / html / csv ─────────────────────────────────

def _read_text(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def parse_text(path: str, file_type: str) -> ParsedDocument:
    text = _read_text(path)
    if file_type == "html":
        text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", text, flags=re.DOTALL | re.IGNORECASE)
        text = re.sub(r"<h([1-6])[^>]*>(.*?)</h\1>", lambda m: "\n" + "#" * int(m.group(1)) + " " + m.group(2) + "\n", text, flags=re.DOTALL | re.IGNORECASE)
        text = re.sub(r"<br\s*/?>|</p>|</div>|</li>|</tr>", "\n", text, flags=re.IGNORECASE)
        text = re.sub(r"<[^>]+>", " ", text)
        text = html.unescape(text)
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
    elif file_type == "csv":
        rows = list(csv.reader(text.splitlines()))
        if rows:
            header = "| " + " | ".join(rows[0]) + " |"
            sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
            body = ["| " + " | ".join(r) + " |" for r in rows[1:]]
            text = "\n".join([header, sep] + body)
    # Paginate long text (~3000 chars) so citations have a locator.
    pages: List[ParsedPage] = []
    page_size = 3000
    blocks = re.split(r"\n\s*\n", text.strip())
    buf: List[str] = []
    size = 0
    for block in blocks:
        if size + len(block) > page_size and buf:
            pages.append(ParsedPage(page=len(pages) + 1, text="\n\n".join(buf)))
            buf, size = [], 0
        buf.append(block)
        size += len(block)
    if buf:
        pages.append(ParsedPage(page=len(pages) + 1, text="\n\n".join(buf)))
    for p in pages:
        p.tables = _extract_md_tables(p.text)
        p.headings = _extract_headings(p.text)
    first_heading = next((h for p in pages for h in p.headings), None)
    return ParsedDocument(filename=os.path.basename(path), file_type=file_type, pages=pages, parser="text",
                          title=first_heading, metadata={"approx_pages": len(pages) > 1})


# ── Entry point ────────────────────────────────────────────────────────

def parse_file(path: str, file_type: Optional[str] = None) -> ParsedDocument:
    if not os.path.exists(path):
        raise ParserError(f"File not found: {path}")
    ext = (file_type or path.rsplit(".", 1)[-1]).lower()
    if ext == "pdf":
        doc = parse_pdf(path)
    elif ext == "docx":
        doc = parse_docx(path)
    elif ext in ("txt", "md", "markdown", "html", "htm", "csv"):
        doc = parse_text(path, "html" if ext in ("html", "htm") else ("md" if ext == "markdown" else ext))
    else:
        raise ParserError(f"Unsupported file type: {ext}")
    if not doc.pages or not any(p.text.strip() or p.tables for p in doc.pages):
        raise ParserError(f"No extractable text in {os.path.basename(path)} (scanned PDF? enable OCR/Docling)")
    logger.info("Parsed document", file=doc.filename, parser=doc.parser, pages=doc.total_pages)
    return doc
