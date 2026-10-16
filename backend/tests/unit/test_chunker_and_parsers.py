from pathlib import Path

import pytest

from src.contracts import ContentType
from src.ingestion import HierarchicalChunker, parse_file
from src.ingestion.parsers import ParserError

SAMPLES = Path(__file__).parent.parent / "samples"


def test_parse_markdown_extracts_title_tables_and_headings():
    doc = parse_file(str(SAMPLES / "policy.md"))
    assert doc.title == "Acme Corp Travel and Expense Policy"
    assert doc.total_pages == 1
    assert len(doc.pages[0].tables) == 1
    assert "Reimbursement Limits" in " ".join(doc.pages[0].headings)


def test_parse_docx_roundtrip(tmp_path):
    from docx import Document
    d = Document()
    d.add_heading("Quarterly Report", level=1)
    d.add_paragraph("Revenue grew 10% in Q3 compared to Q2.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Quarter", "Revenue"
    t.cell(1, 0).text, t.cell(1, 1).text = "Q3", "120"
    path = tmp_path / "r.docx"
    d.save(path)
    doc = parse_file(str(path))
    assert doc.parser == "python-docx"
    assert "# Quarterly Report" in doc.text
    assert "| Quarter | Revenue |" in doc.text


def test_unsupported_and_empty_files(tmp_path):
    bad = tmp_path / "x.xyz"
    bad.write_text("hi")
    with pytest.raises(ParserError):
        parse_file(str(bad))
    empty = tmp_path / "e.txt"
    empty.write_text("   ")
    with pytest.raises(ParserError):
        parse_file(str(empty))


def test_chunker_tracks_sections_tables_and_contextual_headers():
    doc = parse_file(str(SAMPLES / "policy.md"))
    chunks = HierarchicalChunker(chunk_size=120, chunk_overlap=20, contextual_headers=True).chunk_document(doc, "doc-1")
    assert chunks, "no chunks"
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    tables = [c for c in chunks if c.content_type == ContentType.TABLE]
    assert tables and tables[0].raw_content.startswith("| Expense type")
    assert all(c.content.startswith("Document: Acme Corp Travel and Expense Policy") for c in chunks)
    assert all(c.raw_content and not c.raw_content.startswith("Document:") for c in chunks)
    assert all(c.token_estimate <= 120 + 60 for c in chunks if c.content_type == ContentType.TEXT)


def test_chunker_without_headers_and_long_paragraph():
    from src.contracts import ParsedDocument, ParsedPage
    long = "This is sentence number %d of a very long paragraph. " * 1
    text = " ".join(f"This is sentence number {i} of a very long paragraph." for i in range(200))
    doc = ParsedDocument(filename="long.txt", file_type="txt", pages=[ParsedPage(page=1, text=text)])
    chunks = HierarchicalChunker(chunk_size=100, chunk_overlap=10, contextual_headers=False).chunk_document(doc, "d")
    assert len(chunks) > 5
    assert all(c.content == c.raw_content for c in chunks)
    # overlap: consecutive chunks share text
    assert chunks[1].raw_content.split(".")[0].strip() in chunks[0].raw_content
